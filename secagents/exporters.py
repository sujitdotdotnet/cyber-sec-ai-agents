"""Export the assessment as Word (.docx), Excel (.xlsx) and PDF.

All three are built from the same data as the Markdown report (``report.py``) so the numbers and
sections match. Report text comes from logs, code and LLMs, so it is untrusted: control characters are
stripped (they make Word files invalid), PDF markup is escaped, and Excel text that starts with ``=`` is
stored as a string so it can never run as a formula.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any

from xml.sax.saxutils import escape as xml_escape

from .version import version_label

MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

# colours (hex, no #) for severity / status cells
SEVERITY_FILL = {
    "CRITICAL": "C0392B", "HIGH": "E67E22", "MEDIUM": "F4D03F", "LOW": "5DADE2", "INFO": "BDC3C7",
    "FAIL": "C0392B", "PARTIAL": "F4D03F", "PASS": "58D68D",
}
DARK_TEXT_ON = {"MEDIUM", "LOW", "INFO", "PARTIAL", "PASS"}  # light fills need dark text
HEADER_FILL, HEADER_TEXT = "1F3A5F", "FFFFFF"

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]")


def clean(value: Any) -> Any:
    """Strip characters that are invalid in XML/Office files; other values pass through."""
    return _CONTROL.sub("", value) if isinstance(value, str) else value


def text(value: Any) -> str:
    return "" if value is None else str(clean(value))


# ---------------------------------------------------------------------------------------------
# data model shared by the three formats
# ---------------------------------------------------------------------------------------------
@dataclass
class Table:
    title: str
    header: list[str]
    rows: list[list[Any]]
    sev_cols: tuple[int, ...] = ()  # columns coloured by severity / status


@dataclass
class Block:
    kind: str  # title | subtitle | h2 | h3 | p | bullets | numbered | code | table | gap
    text: str = ""
    items: list[str] = field(default_factory=list)
    table: Table | None = None


def _sev(rows: list[list[Any]], col: int) -> None:  # normalise to upper-case so colours always match
    for r in rows:
        r[col] = text(r[col]).upper()


def datasets(state: dict[str, Any], full: bool = False) -> dict[str, Table]:
    """The report tables. ``full=True`` adds the extra columns/tables used by the Excel workbook."""
    logs = state.get("log_findings", {}) or {}
    vulns = state.get("vuln_findings", {}) or {}
    intel = state.get("threat_intel", {}) or {}
    ir = state.get("incident_plan", {}) or {}
    comp = state.get("compliance", {}) or {}
    t: dict[str, Table] = {}

    ev = logs.get("events", [])
    if full:
        t["Detections"] = Table("Log detections", ["ID", "Severity", "Detection", "Source file", "Source IP", "Target", "Count",
                                                    "First seen", "Last seen", "ATT&CK", "Description", "Evidence"], [
            [e["id"], e["severity"], e["type"], e.get("source", ""), e.get("src_ip", ""), e.get("target", ""), e.get("count", ""),
             e.get("first_seen", ""), e.get("last_seen", ""), ", ".join(e.get("mitre", [])), e.get("description", ""),
             "\n".join(e.get("evidence", []))] for e in ev], (1,))
    else:
        t["Detections"] = Table("Log detections", ["ID", "Severity", "Detection", "Source", "Count", "ATT&CK"], [
            [e["id"], e["severity"], e["type"], e.get("src_ip") or e.get("target", ""), e.get("count", ""),
             ", ".join(e.get("mitre", []))] for e in ev], (1,))

    t["CVEs"] = Table("CVEs affecting deployed components", ["CVE", "Component", "Installed", "Fixed in", "CVSS", "Severity", "Note", "Confidence", "Data source", "CISA KEV"], [
        [m["cve"], m["package"], m["installed"], m["fixed"], m["cvss"], m.get("severity", ""), m.get("note", ""), m.get("confidence", ""),
         m.get("data_source", ""), "exploited" if m.get("exploited_in_the_wild") else ""]
        for m in intel.get("cve_matches", [])], (5,))
    t["Attack correlation"] = Table("Attacks that hit real weaknesses", ["Attack", "Source", "Code location", "Severity", "Verdict", "Confidence"], [
        [c["attack"], c["src_ip"], c["location"], c.get("severity", ""), c["verdict"], c.get("confidence", "")] for c in intel.get("correlations", [])], (3,))
    t["Exploit attempts"] = Table("Exploit attempts", ["Attack", "CVE", "Component", "Present", "Verdict", "Confidence"], [
        [a["attack"], a["cve"], a.get("component", ""), "yes" if a.get("component_present") else "no", a["verdict"], a.get("confidence", "")]
        for a in intel.get("exploit_attempts", [])])
    t["IOCs"] = Table("Known-bad indicators", ["Indicator", "Reputation", "Tags", "Seen in"], [
        [i["indicator"], i["reputation"], ", ".join(i["tags"]), i.get("seen_in", "")] for i in intel.get("ioc_hits", [])])

    fi = vulns.get("findings", [])
    if full:
        t["Findings"] = Table("Code, config and dependency findings", ["ID", "Severity", "Category", "Finding", "File", "Line", "CWE", "Tool",
                                                                       "Rule", "Snippet"], [
            [f["id"], f["severity"], f.get("category", ""), f["title"], f.get("file", ""), f.get("line", ""), f.get("cwe", ""),
             f.get("tool", ""), f.get("rule", ""), f.get("snippet", "")] for f in fi], (1,))
    else:
        t["Findings"] = Table("Code, config and dependency findings", ["ID", "Severity", "Category", "Finding", "Location", "CWE"], [
            [f["id"], f["severity"], f.get("category", ""), f["title"], f"{f.get('file', '')}:{f.get('line', '')}".rstrip(":"),
             f.get("cwe", "")] for f in fi], (1,))

    t["Incidents"] = Table("Incidents", ["Incident", "Severity", "Priority", "Title", "Summary"], [
        [i["id"], i["severity"], i["priority"], i["title"], i["summary"]] for i in ir.get("incidents", [])], (1,))
    t["Playbook"] = Table("Incident response playbook (NIST SP 800-61)", ["#", "Phase", "Priority", "Action", "Command", "Owner"], [
        [s["step"], s["phase"], s["priority"], s["action"], s.get("command", ""), s["owner"]] for s in ir.get("steps", [])])
    corr, critic, gate = state.get("correlation", {}) or {}, state.get("critic", {}) or {}, state.get("containment_gate", {}) or {}
    t["Correlation"] = Table("Correlated incidents (events, weaknesses, CVEs, techniques, control gaps)", ["Incident", "Title", "Events", "Attackers", "Code weaknesses", "CVEs", "ATT&CK", "Known-bad", "Control gaps exposed"], [
        [ln["incident"], ln["title"], ", ".join(ln["events"]), ", ".join(ln["attackers"]), ", ".join(ln["weaknesses"]), ", ".join(ln["cves"]),
         ", ".join(ln["techniques"]), ", ".join(ln["known_bad"]), "; ".join(ln["control_gaps"])] for ln in corr.get("links", [])])
    t["Containment"] = Table("Containment and eradication actions (human approval gate; nothing is executed by the tool)", ["ID", "Incident", "Phase", "Priority", "Action", "Proposed command", "Status", "Decided by", "Approved command"], [
        [a["id"], a["incident"], a["phase"], a["priority"], a["action"], a["command"], a["status"], a["decided_by"], a["approved_command"]] for a in gate.get("actions", [])])
    t["Approval audit"] = Table("Approval audit trail", ["When (UTC)", "By", "Action", "Decision", "Command before", "Command after"], [
        [e["at"], e["by"], e["action"], e["decision"], e["command_before"], e["command_after"]] for e in gate.get("audit", [])])
    t["Review"] = Table("Uncertain findings (uncertainty critic)", ["Kind", "Item", "Confidence", "Why"], [
        [u["kind"], u["id"], u["confidence"], u["reason"]] for u in critic.get("uncertain", [])])
    t["Unverified"] = Table("Exposures that could not be verified", ["Component", "Ecosystem", "Why"], [
        [u["component"], u["ecosystem"], u["reason"]] for u in critic.get("unverified", [])])
    t["Compliance"] = Table("Framework scores", ["Framework", "Score %", "Pass", "Partial", "Fail", "Controls"], [
        [f["framework"], f["score"], f["pass"], f["partial"], f["fail"], f["total"]] for f in comp.get("frameworks", [])])
    t["Fixes"] = Table("Prioritized fixes (by controls closed)", ["Fix", "Controls closed", "Evidence", "Recommended action"], [
        [g["title"], g["controls_affected"], g["evidence"], g["fix"]] for g in comp.get("gaps", [])])
    t["Control gaps"] = Table("Controls that are not passing", ["Framework", "Control", "Title", "Status", "Evidence"], [
        [c["framework"], c["id"], c["title"], c["status"], c["evidence"]] for c in comp.get("controls", []) if c["status"] != "PASS"], (3,))

    if full:
        t["Checks"] = Table("Technical checks", ["Check", "Result", "Evidence", "Fix"], [
            [c["title"], "PASS" if c["passed"] else "FAIL", c["evidence"], c["fix"]] for c in comp.get("checks", [])], (1,))
        t["Inventory"] = Table("Software inventory", ["Name", "Version", "Ecosystem", "Source"], [
            [i["name"], i["version"], i["ecosystem"], i["source"]] for i in vulns.get("inventory", [])])
    return t


def summary_rows(state: dict[str, Any]) -> list[list[Any]]:
    logs = state.get("log_findings", {}) or {}
    vulns = state.get("vuln_findings", {}) or {}
    intel = state.get("threat_intel", {}) or {}
    ir = state.get("incident_plan", {}) or {}
    comp = state.get("compliance", {}) or {}
    rows: list[list[Any]] = [
        ["Overall risk", intel.get("overall_risk", "n/a")],
        ["Incidents", len(ir.get("incidents", []))],
        ["Highest incident priority", ir.get("highest_priority", "n/a")],
        ["Log detections", len(logs.get("events", []))],
        ["Code / config findings", len(vulns.get("findings", []))],
        ["CVEs affecting inventory", len(intel.get("cve_matches", []))],
        ["Attacker IPs", ", ".join(logs.get("attackers", [])) or "none"],
    ]
    rows += [[f"Compliance: {f['framework']}", f"{f['score']}%"] for f in comp.get("frameworks", [])]
    critic, gate = state.get("critic", {}) or {}, state.get("containment_gate", {}) or {}
    if critic:
        c = critic.get("counts", {})
        rows.append(["Findings reviewed by the critic (high / medium / low confidence)", f"{c.get('high', 0)} / {c.get('medium', 0)} / {c.get('low', 0)}"])
    if gate.get("actions"):
        n = {s: sum(1 for a in gate["actions"] if a["status"] == s) for s in ("pending", "approved", "rejected")}
        rows.append(["Containment actions (this tool executes none of them)", f"{n['pending']} awaiting approval, {n['approved']} approved, {n['rejected']} rejected"])
    if state.get("redactions"):
        rows.append(["Secrets / personal data removed before reaching a model", state["redactions"]])
    return rows


def narratives(state: dict[str, Any]) -> list[tuple[str, str]]:
    """The markdown text sections, in report order."""
    return [
        ("Incident briefing", (state.get("incident_plan") or {}).get("briefing", "")),
        ("Log Monitor: what happened", (state.get("log_findings") or {}).get("narrative", "")),
        ("Threat Intelligence: are we affected?", (state.get("threat_intel") or {}).get("summary", "")),
        ("Vulnerability Scanner: suggested fixes", (state.get("vuln_findings") or {}).get("remediation", "")),
        ("Policy Checker: remediation roadmap", (state.get("compliance") or {}).get("roadmap", "")),
    ]


# ---------------------------------------------------------------------------------------------
# tiny markdown reader for the LLM/narrative text (headings, bullets, numbered lists, code, tables, **bold**)
# ---------------------------------------------------------------------------------------------
_INLINE = re.compile(r"(\*\*(?P<b>.+?)\*\*|`(?P<c>[^`]+)`|(?<![\w*])\*(?P<i1>[^*\s][^*]*?)\*(?![\w*])|(?<![\w_])_(?P<i2>[^_\s][^_]*?)_(?![\w_]))")


def inline_runs(s: str) -> list[tuple[str, str]]:
    """Split text into (text, style) runs where style is '', 'b', 'i' or 'c' (code)."""
    runs, pos = [], 0
    for m in _INLINE.finditer(s):
        if m.start() > pos:
            runs.append((s[pos:m.start()], ""))
        if m.group("b") is not None:
            runs.append((m.group("b"), "b"))
        elif m.group("c") is not None:
            runs.append((m.group("c"), "c"))
        else:
            runs.append((m.group("i1") or m.group("i2"), "i"))
        pos = m.end()
    if pos < len(s):
        runs.append((s[pos:], ""))
    return runs or [("", "")]


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def md_blocks(md: str) -> list[Block]:
    blocks: list[Block] = []
    lines = text(md).replace("\r\n", "\n").split("\n")
    para: list[str] = []
    bullets: list[str] = []
    numbered: list[str] = []
    i = 0

    def flush() -> None:
        nonlocal para, bullets, numbered
        if para:
            blocks.append(Block("p", " ".join(para)))
        if bullets:
            blocks.append(Block("bullets", items=bullets))
        if numbered:
            blocks.append(Block("numbered", items=numbered))
        para, bullets, numbered = [], [], []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append(Block("code", "\n".join(code)))
        elif stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", lines[i + 1]) and "-" in lines[i + 1]:
            flush()
            header = _split_row(stripped)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(_split_row(lines[i]))
                i += 1
            i -= 1
            blocks.append(Block("table", table=Table("", header, [r + [""] * (len(header) - len(r)) for r in rows])))
        elif re.match(r"^#{1,6}\s+", stripped):
            flush()
            blocks.append(Block("h3", re.sub(r"^#{1,6}\s+", "", stripped)))
        elif re.match(r"^[-*•]\s+", stripped):
            if para or numbered:
                flush()
            bullets.append(re.sub(r"^[-*•]\s+", "", stripped))
        elif re.match(r"^\d+[.)]\s+", stripped):
            if para or bullets:
                flush()
            numbered.append(re.sub(r"^\d+[.)]\s+", "", stripped))
        elif not stripped:
            flush()
        else:
            if bullets or numbered:  # a wrapped continuation line of the last list item
                (bullets or numbered)[-1] += " " + stripped
            else:
                para.append(stripped)
        i += 1
    flush()
    return blocks


def build_blocks(state: dict[str, Any], meta: dict[str, Any] | None = None) -> list[Block]:
    meta = meta or {}
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    t = datasets(state)
    nar = dict(narratives(state))
    sub = f"Generated {now} by SecAgents {version_label()}."
    if meta.get("name"):
        sub = f"{meta['name']}. {sub}"
    b: list[Block] = [Block("title", "Security Assessment Report"), Block("subtitle", sub), Block("h2", "Executive summary"),
                      Block("table", table=Table("", ["Metric", "Value"], summary_rows(state), sev_cols=(1,)))]
    b += md_blocks(nar["Incident briefing"])

    def section(title: str, keys: list[str], intro: str = "") -> None:
        b.append(Block("h2", title))
        b.extend(md_blocks(intro))
        for k in keys:
            tb = t[k]
            b.append(Block("h3", tb.title))
            b.append(Block("table", table=tb) if tb.rows else Block("p", "None."))

    section("1. Log Monitor: what happened", ["Detections"], nar["Log Monitor: what happened"])
    section("2. Threat Intelligence: are we affected?", ["CVEs", "Attack correlation", "Exploit attempts", "IOCs"],
            nar["Threat Intelligence: are we affected?"])
    section("3. Vulnerability Scanner", ["Findings"])
    b.append(Block("h3", "Suggested fixes"))
    b.extend(md_blocks(nar["Vulnerability Scanner: suggested fixes"]))
    section("4. Incident Response playbook", ["Incidents", "Correlation", "Playbook", "Containment"] + (["Approval audit"] if t["Approval audit"].rows else []))
    section("5. Policy Checker: compliance gaps", ["Compliance"], "")
    b.extend(md_blocks(nar["Policy Checker: remediation roadmap"]))
    for k in ("Fixes", "Control gaps"):
        b.append(Block("h3", t[k].title))
        b.append(Block("table", table=t[k]) if t[k].rows else Block("p", "None."))
    section("6. Uncertainty review: what to treat with caution", ["Review", "Unverified"],
            (state.get("critic", {}) or {}).get("caveat", "") or "The critic found nothing that needs a caveat.")
    return b


def _weights(tb: Table) -> list[float]:
    """Relative column widths from content length: short ids stay narrow, long text gets room."""
    w = []
    for c, h in enumerate(tb.header):
        longest = max([len(text(h))] + [min(len(text(r[c])), 60) for r in tb.rows if c < len(r)])
        w.append(max(5, min(longest, 60)))
    return w


def _fill_for(value: Any) -> str | None:
    return SEVERITY_FILL.get(text(value).strip().upper())


# ---------------------------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------------------------
def to_docx(state: dict[str, Any], meta: dict[str, Any] | None = None) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    usable_cm = 29.7 - 3.2
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.core_properties.title = "Security Assessment Report"
    doc.core_properties.author = "SecAgents"

    def shade(cell, hex_fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc_pr.append(shd)

    def add_runs(par, s: str, size: float | None = None, color: str | None = None, bold: bool = False) -> None:
        for chunk, style in inline_runs(text(s)):
            run = par.add_run(chunk)
            run.bold = bold or style == "b"
            run.italic = style == "i"
            if style == "c":
                run.font.name, run.font.size = "Consolas", Pt((size or 10) - 1)
            elif size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)

    for blk in build_blocks(state, meta):
        if blk.kind == "title":
            p = doc.add_paragraph()
            add_runs(p, blk.text, size=26, color=HEADER_FILL, bold=True)
        elif blk.kind == "subtitle":
            p = doc.add_paragraph()
            add_runs(p, blk.text, size=10, color="666666")
        elif blk.kind in ("h2", "h3"):
            doc.add_heading(text(blk.text), level=1 if blk.kind == "h2" else 2)
        elif blk.kind == "p":
            add_runs(doc.add_paragraph(), blk.text)
        elif blk.kind == "bullets":
            for it in blk.items:
                add_runs(doc.add_paragraph(style="List Bullet"), it)
        elif blk.kind == "numbered":
            for n, it in enumerate(blk.items, 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                add_runs(p, f"{n}. {it}")
        elif blk.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(text(blk.text))
            run.font.name, run.font.size = "Consolas", Pt(8.5)
        elif blk.kind == "table" and blk.table is not None:
            tb = blk.table
            n = len(tb.header)
            table = doc.add_table(rows=1, cols=n)
            table.style = "Table Grid"
            table.autofit = False
            weights = _weights(tb)
            widths = [Cm(usable_cm * w / sum(weights)) for w in weights]
            for c, h in enumerate(tb.header):
                cell = table.rows[0].cells[c]
                cell.width = widths[c]
                shade(cell, HEADER_FILL)
                add_runs(cell.paragraphs[0], h, size=8.5, color=HEADER_TEXT, bold=True)
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)  # header row repeats on every page
            for row in tb.rows:
                cells = table.add_row().cells
                for c in range(n):
                    value = row[c] if c < len(row) else ""
                    cells[c].width = widths[c]
                    fill = _fill_for(value) if c in tb.sev_cols else None
                    if fill:
                        shade(cells[c], fill)
                    dark = text(value).upper() in DARK_TEXT_ON
                    add_runs(cells[c].paragraphs[0], text(value)[:1500], size=8.5,
                             color=("FFFFFF" if fill and not dark else None), bold=bool(fill))
            doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------------------------
_PDF_SUBS = {"→": "->", "←": "<-", "≥": ">=", "≤": "<=", "✓": "OK", "✗": "x", "…": "...", " ": " ", "‑": "-"}


def _pdf_text(value: Any) -> str:
    """Standard PDF fonts cover Latin-1 only: map common symbols, drop what cannot be shown (emoji etc.)."""
    s = text(value)
    for a, b in _PDF_SUBS.items():
        s = s.replace(a, b)
    return s.encode("cp1252", "ignore").decode("cp1252")


def _pdf_markup(value: Any, size: float | None = None) -> str:
    out = []
    for chunk, style in inline_runs(_pdf_text(value)):
        esc = xml_escape(chunk)
        if style == "b":
            esc = f"<b>{esc}</b>"
        elif style == "i":
            esc = f"<i>{esc}</i>"
        elif style == "c":
            esc = f'<font face="Courier" size="{(size or 9.5) - 1}">{esc}</font>'
        out.append(esc)
    return "".join(out).replace("\n", "<br/>")


def to_pdf(state: dict[str, Any], meta: dict[str, Any] | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import CondPageBreak, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table as RLTable, TableStyle

    page = landscape(A4)
    margin = 1.5 * cm
    usable = page[0] - 2 * margin
    base = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5)
    h2 = ParagraphStyle("h2", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#" + HEADER_FILL),
                        spaceBefore=12, spaceAfter=6)
    h3 = ParagraphStyle("h3", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, spaceBefore=8, spaceAfter=4)
    title = ParagraphStyle("title", parent=base["Title"], fontName="Helvetica-Bold", fontSize=26, leading=32, alignment=0,
                           textColor=colors.HexColor("#" + HEADER_FILL), spaceAfter=4)
    subtitle = ParagraphStyle("sub", parent=body, textColor=colors.HexColor("#666666"), spaceAfter=10)
    code = ParagraphStyle("code", parent=body, fontName="Courier", fontSize=8, leading=10, backColor=colors.HexColor("#F2F3F4"),
                          borderPadding=4, wordWrap="CJK")
    cell = ParagraphStyle("cell", parent=body, fontSize=7.6, leading=9.4, spaceAfter=0, wordWrap="CJK")
    cell_h = ParagraphStyle("cellh", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)

    story: list[Any] = []
    for blk in build_blocks(state, meta):
        if blk.kind == "title":
            story.append(Paragraph(_pdf_markup(blk.text), title))
        elif blk.kind == "subtitle":
            story.append(Paragraph(_pdf_markup(blk.text), subtitle))
        elif blk.kind in ("h2", "h3"):
            story.append(CondPageBreak(3.5 * cm))  # a heading never ends a page: break if little room is left
            story.append(Paragraph(_pdf_markup(blk.text), h2 if blk.kind == "h2" else h3))
        elif blk.kind == "p":
            story.append(Paragraph(_pdf_markup(blk.text), body))
        elif blk.kind in ("bullets", "numbered"):
            items = [ListItem(Paragraph(_pdf_markup(it), body), leftIndent=14) for it in blk.items]
            story.append(ListFlowable(items, bulletType="bullet" if blk.kind == "bullets" else "1", start=None if blk.kind == "bullets" else 1,
                                      bulletFontSize=8, leftIndent=14))
        elif blk.kind == "code":
            story.append(Paragraph(_pdf_markup(blk.text).replace(" ", "&nbsp;"), code))
        elif blk.kind == "table" and blk.table is not None:
            tb = blk.table
            weights = _weights(tb)
            widths = [usable * w / sum(weights) for w in weights]
            data = [[Paragraph(_pdf_markup(h), cell_h) for h in tb.header]]
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HEADER_FILL)),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BFC5CC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ]
            for r, row in enumerate(tb.rows, 1):
                line = []
                for c in range(len(tb.header)):
                    value = row[c] if c < len(row) else ""
                    fill = _fill_for(value) if c in tb.sev_cols else None
                    st_cell = cell
                    if fill:
                        style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor("#" + fill)))
                        dark = text(value).upper() in DARK_TEXT_ON
                        st_cell = ParagraphStyle("sev", parent=cell, fontName="Helvetica-Bold", textColor=colors.black if dark else colors.white)
                    line.append(Paragraph(_pdf_markup(text(value)[:1200]), st_cell))
                data.append(line)
            story.append(RLTable(data, colWidths=widths, repeatRows=1, style=TableStyle(style)))
            story.append(Spacer(1, 8))

    def footer(canvas, doc_) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawString(margin, 0.9 * cm, f"SecAgents {version_label()} security assessment")
        canvas.drawRightString(page[0] - margin, 0.9 * cm, f"Page {doc_.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=page, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=1.6 * cm,
                      title="Security Assessment Report", author="SecAgents").build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------------------------
def to_xlsx(state: dict[str, Any], meta: dict[str, Any] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    meta = meta or {}
    wb = Workbook()
    head_font = Font(bold=True, color=HEADER_TEXT)
    head_fill = PatternFill("solid", fgColor=HEADER_FILL)
    wrap = Alignment(wrap_text=True, vertical="top")

    def put(ws, row: int, col: int, value: Any):
        value = clean(value)
        c = ws.cell(row=row, column=col, value=value)
        if isinstance(value, str) and value.startswith("="):
            c.data_type = "s"  # log / code text must never become a formula
        return c

    def style_fill(c, raw: Any) -> None:
        fill = _fill_for(raw)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
            c.font = Font(bold=True, color="000000" if text(raw).upper() in DARK_TEXT_ON else "FFFFFF")

    def add_sheet(name: str, tb: Table, first: bool = False) -> None:
        ws = wb.active if first else wb.create_sheet()
        ws.title = name[:31]
        for c, h in enumerate(tb.header, 1):
            cell = put(ws, 1, c, h)
            cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center")
        for r, row in enumerate(tb.rows, 2):
            for c, value in enumerate(row, 1):
                cell = put(ws, r, c, value)
                cell.alignment = wrap
                if (c - 1) in tb.sev_cols:
                    style_fill(cell, value)
        for c, w in enumerate(_weights(tb), 1):
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 8), 70)
        ws.freeze_panes = "A2"
        if tb.rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(tb.header))}{len(tb.rows) + 1}"

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws["A1"], ws["A1"].font = "Security Assessment Report", Font(bold=True, size=16, color=HEADER_FILL)
    ws["A2"] = f"{meta.get('name') + '. ' if meta.get('name') else ''}Generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M UTC} by SecAgents {version_label()}."
    for c, h in enumerate(("Metric", "Value"), 1):
        cell = put(ws, 4, c, h)
        cell.font, cell.fill = head_font, head_fill
    for r, (k, v) in enumerate(summary_rows(state), 5):
        put(ws, r, 1, k).font = Font(bold=True)
        cell = put(ws, r, 2, v)
        cell.alignment = Alignment(horizontal="left", wrap_text=True)
        if k == "Overall risk":
            style_fill(cell, v)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 70

    tables = datasets(state, full=True)
    for name, tb in tables.items():
        add_sheet(name, tb)

    nar = wb.create_sheet("Narrative")
    for c, h in enumerate(("Section", "Text"), 1):
        cell = put(nar, 1, c, h)
        cell.font, cell.fill = head_font, head_fill
    for r, (title_, body) in enumerate(narratives(state), 2):
        put(nar, r, 1, title_).font = Font(bold=True)
        cell = put(nar, r, 2, body)
        cell.alignment = wrap
    nar.column_dimensions["A"].width, nar.column_dimensions["B"].width = 38, 130
    nar.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


EXPORTERS = {"docx": to_docx, "xlsx": to_xlsx, "pdf": to_pdf}


def export(fmt: str, state: dict[str, Any], meta: dict[str, Any] | None = None) -> bytes:
    return EXPORTERS[fmt](state, meta)
