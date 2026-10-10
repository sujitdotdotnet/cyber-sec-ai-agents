import io
import zipfile

import pytest

from secagents import exporters
from secagents.graph import run


@pytest.fixture(scope="module")
def state():
    return run()  # the full demo pipeline, rule-based (conftest forces offline)


def test_all_formats_are_valid_files(state):
    docx = exporters.export("docx", state, {"name": "t"})
    xlsx = exporters.export("xlsx", state, {"name": "t"})
    pdf = exporters.export("pdf", state, {"name": "t"})
    assert pdf.startswith(b"%PDF") and pdf.rstrip().endswith(b"%%EOF")
    for blob, member in ((docx, "word/document.xml"), (xlsx, "xl/workbook.xml")):
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            assert z.testzip() is None and member in z.namelist()


def test_docx_matches_the_report_data(state):
    from docx import Document

    doc = Document(io.BytesIO(exporters.to_docx(state, {})))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads[:2] == ["Executive summary", "1. Log Monitor: what happened"]
    assert "5. Policy Checker: compliance gaps" in heads
    id_tables = [t for t in doc.tables if [c.text for c in t.rows[0].cells[:2]] == ["ID", "Severity"]]
    detections, findings = id_tables[0], id_tables[1]  # report order: log detections, then scanner findings
    assert len(detections.rows) - 1 == len(state["log_findings"]["events"])
    assert len(findings.rows) - 1 == len(state["vuln_findings"]["findings"])
    assert detections.rows[1].cells[1]._tc.tcPr.xpath("./w:shd/@w:fill") == [exporters.SEVERITY_FILL["CRITICAL"]]


def test_xlsx_has_a_sheet_per_table_with_real_numbers(state):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(exporters.to_xlsx(state, {})))
    assert {"Summary", "Detections", "CVEs", "Findings", "Playbook", "Compliance", "Control gaps", "Narrative"} <= set(wb.sheetnames)
    assert wb["Detections"].max_row - 1 == len(state["log_findings"]["events"])
    assert wb["Findings"].max_row - 1 == len(state["vuln_findings"]["findings"])
    assert isinstance(wb["CVEs"]["E2"].value, float)  # CVSS stays numeric, so it can be sorted and filtered
    assert wb["Detections"].freeze_panes == "A2" and wb["Detections"].auto_filter.ref


def test_hostile_content_cannot_break_or_attack_the_files(state):
    import copy

    from openpyxl import load_workbook

    bad = copy.deepcopy(state)
    ev = bad["log_findings"]["events"][0]
    ev["evidence"] = ['=HYPERLINK("http://evil.example","click")', "x\x00y\x07z"]
    ev["description"] = "<script>alert(1)</script> & <b>unclosed </para> \U0001f916 \u2192"
    bad["incident_plan"]["steps"][0]["command"] = "+cmd|' /C calc'!A0"
    bad["log_findings"]["narrative"] = "**bold** <not-a-tag> `code` \x01 end"
    for fmt in ("docx", "xlsx", "pdf"):
        assert exporters.export(fmt, bad, {})  # control chars, stray markup and emoji must not raise
    cell = load_workbook(io.BytesIO(exporters.to_xlsx(bad, {})))["Detections"]["L2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")  # text, never a formula
    assert "\x00" not in cell.value


def test_markdown_text_is_parsed_into_structure():
    blocks = exporters.md_blocks("Intro **bold**\n\n- one\n- two\n  wrapped\n\n1. a\n2. b\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n```\ncode\n```")
    assert [b.kind for b in blocks] == ["p", "bullets", "numbered", "table", "code"]
    assert blocks[1].items == ["one", "two wrapped"]
    assert blocks[3].table.rows == [["1", "2"]]
    assert exporters.inline_runs("a **b** `c` d") == [("a ", ""), ("b", "b"), (" ", ""), ("c", "c"), (" d", "")]


def test_cli_writes_the_new_formats(tmp_path, monkeypatch):
    import sys

    from secagents.__main__ import main

    out = {f: tmp_path / f"r.{f}" for f in ("docx", "xlsx", "pdf")}
    monkeypatch.setattr(sys, "argv", ["secagents", "--out", str(tmp_path / "r.md")] + [a for f, p in out.items() for a in (f"--{f}", str(p))])
    main()
    assert all(p.stat().st_size > 5000 for p in out.values()) and (tmp_path / "r.md").exists()
