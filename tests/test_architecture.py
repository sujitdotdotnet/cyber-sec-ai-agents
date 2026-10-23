"""The supervisor architecture: parallel specialists, correlation, the uncertainty critic and the human gate."""

import copy
import threading
import time

import pytest

from secagents import approvals, graph
from secagents.agents import AGENTS
from secagents.correlate import correlate
from secagents.critic import review
from secagents.graph import build_graph, eligible, next_agent, plan_waves


# ---- planning and parallel execution -----------------------------------------------------------------------------------
def test_the_plan_runs_independent_specialists_together():
    assert plan_waves() == [["log_monitor", "vuln_scanner"], ["threat_intel"], ["incident_response", "policy_checker"]]
    assert eligible([]) == ["log_monitor", "vuln_scanner"]
    assert eligible(["log_monitor"]) == ["vuln_scanner"]
    assert eligible(["log_monitor", "vuln_scanner", "threat_intel"]) == ["incident_response", "policy_checker"]
    assert eligible(list(graph.DEPENDENCIES)) == []


def test_the_sequential_helper_still_agrees_with_the_plan():
    assert next_agent([]) == "log_monitor" and next_agent(["log_monitor"]) == "vuln_scanner"
    assert next_agent(["log_monitor", "vuln_scanner"]) == "threat_intel" and next_agent(list(graph.DEPENDENCIES)) == "reporter"


def test_a_cycle_in_the_dependencies_is_refused_not_looped_forever(monkeypatch):
    monkeypatch.setattr(graph, "DEPENDENCIES", {"a": ["b"], "b": ["a"]})
    with pytest.raises(ValueError, match="cycle"):
        plan_waves()


@pytest.fixture
def fake_agents(monkeypatch):
    """Every specialist replaced by a stub that records when it ran; pairs wait for each other at a barrier."""
    log, barriers = [], {"first": threading.Barrier(2, timeout=5), "last": threading.Barrier(2, timeout=5)}
    pair = {"log_monitor": "first", "vuln_scanner": "first", "incident_response": "last", "policy_checker": "last"}

    def make(name):
        def run(state):
            log.append(("start", name, time.monotonic()))
            if name in pair:
                barriers[pair[name]].wait()  # only passes if its partner is running at the same time
            log.append(("end", name, time.monotonic()))
            return {"completed": [name], "trace": [], **{
                "log_monitor": {"log_findings": {"events": []}}, "vuln_scanner": {"vuln_findings": {"inventory": [], "findings": []}},
                "threat_intel": {"threat_intel": {"cve_matches": [], "exploit_attempts": [], "correlations": [], "ioc_hits": [], "summary": ""}},
                "incident_response": {"incident_plan": {"incidents": [], "steps": []}}, "policy_checker": {"compliance": {"gaps": [], "frameworks": []}},
            }[name]}
        return run

    for name in list(AGENTS):
        monkeypatch.setitem(AGENTS, name, make(name))
    return log


def test_independent_specialists_really_run_at_the_same_time(fake_agents):
    # if they ran one after the other, the first would wait at the barrier for a partner that never starts, and this would time out
    state = build_graph().invoke({"inputs": {}, "completed": [], "trace": []})
    assert sorted(state["completed"]) == sorted(graph.DEPENDENCIES)
    order = [(kind, name) for kind, name, _ in fake_agents]
    started = {name: i for i, (kind, name) in enumerate(order) if kind == "start"}
    ended = {name: i for i, (kind, name) in enumerate(order) if kind == "end"}
    assert started["threat_intel"] > max(ended["log_monitor"], ended["vuln_scanner"])  # waits for both of its inputs
    assert started["incident_response"] > ended["threat_intel"] and started["policy_checker"] > ended["threat_intel"]


def test_the_supervisor_runs_once_per_wave_and_the_review_steps_follow_in_order(fake_agents):
    state = build_graph().invoke({"inputs": {}, "completed": [], "trace": []})
    msgs = [(t["agent"], t["message"]) for t in state["trace"]]
    dispatches = [m for a, m in msgs if a == "supervisor" and m.startswith("dispatching")]
    assert dispatches == ["dispatching log_monitor + vuln_scanner in parallel", "dispatching threat_intel",
                          "dispatching incident_response + policy_checker in parallel"]
    assert msgs[0][1].startswith("plan: log_monitor + vuln_scanner → threat_intel → incident_response + policy_checker")
    agents_seen = [a for a, _ in msgs]
    tail = [a for a in agents_seen if a in ("correlator", "critic", "human_gate", "reporter")]
    assert tail == ["correlator", "critic", "human_gate", "reporter"]
    assert agents_seen.index("correlator") > max(i for i, a in enumerate(agents_seen) if a == "supervisor") - 1
    assert {"correlation", "critic", "containment_gate", "report_md"} <= set(state)


# ---- the critic -------------------------------------------------------------------------------------------------------------
def cve(**kw):
    base = {"cve": "CVE-2020-1", "package": "pyyaml", "installed": "5.3", "affected_spec": "<5.4", "cvss": 9.8, "severity": "CRITICAL",
            "ecosystem": "pypi", "data_source": "Knowledge base", "exploited_in_the_wild": False}
    return {**base, **kw}


def state_with(cves=(), attempts=(), corr=(), iocs=(), events=(), inventory=(), summary="Model summary.", briefing="Model briefing."):
    return {"threat_intel": {"cve_matches": [cve(**c) if isinstance(c, dict) else c for c in cves], "exploit_attempts": list(attempts),
                             "correlations": list(corr), "ioc_hits": list(iocs), "summary": summary},
            "incident_plan": {"incidents": [], "steps": [], "briefing": briefing}, "log_findings": {"events": list(events)},
            "vuln_findings": {"inventory": list(inventory)}}


def conf(patch, i=0):
    return patch["threat_intel"]["cve_matches"][i]["confidence"]


def test_a_clean_version_matched_cve_is_high_confidence_and_says_how_it_was_checked():
    patch = review(state_with(cves=[{"exploited_in_the_wild": True}]))
    m = patch["threat_intel"]["cve_matches"][0]
    assert m["confidence"] == "high" and "re-checked against the affected range" in m["confidence_reason"] and "CISA" in m["confidence_reason"]
    assert patch["critic"]["caveat"] == "" and "Review notes" not in patch["threat_intel"]["summary"]  # nothing doubtful, nothing added


def test_the_critic_rechecks_the_version_itself_and_distrusts_a_mismatch():
    patch = review(state_with(cves=[{"installed": "5.4", "affected_spec": "<5.4"}]))  # 5.4 is NOT affected
    assert conf(patch) == "low" and "NOT inside the affected range" in patch["threat_intel"]["cve_matches"][0]["confidence_reason"]
    ok = review(state_with(cves=[{"installed": "1.0", "affected_spec": ">=0.5,<1.2 || >=3.0"}]))
    assert conf(ok) == "high"


def test_distribution_packages_and_unscored_cves_are_medium_confidence():
    assert conf(review(state_with(cves=[{"ecosystem": "os", "package": "openssl"}]))) == "medium"
    p = review(state_with(cves=[{"cvss": 0.0, "severity": "UNKNOWN"}]))
    assert conf(p) == "medium" and "no CVSS score" in p["threat_intel"]["cve_matches"][0]["confidence_reason"]


def test_live_osv_matches_without_a_spec_are_explained():
    p = review(state_with(cves=[{"affected_spec": "", "data_source": "OSV (live)"}]))
    assert conf(p) == "high" and "matched by OSV for this exact version" in p["threat_intel"]["cve_matches"][0]["confidence_reason"]


def test_exploit_attempts_correlations_and_indicators_are_graded():
    attempts = [{"event": "LOG-1", "cve": "CVE-2021-44228", "component": "log4j-core", "component_present": False},
                {"event": "LOG-2", "cve": "CVE-2020-1", "component": "pyyaml", "component_present": True}]
    corr = [{"event": "LOG-3", "finding": "VULN-1", "verdict": "Likely exploited: attack returned HTTP 2xx and the code is vulnerable"},
            {"event": "LOG-4", "finding": "VULN-2", "verdict": "Exploitable: matching weakness exists in code"}]
    iocs = [{"indicator": "1.2.3.4", "reputation": "malicious", "source": "feed"}, {"indicator": "5.6.7.8", "reputation": "suspicious", "source": "feed"}]
    t = review(state_with(attempts=attempts, corr=corr, iocs=iocs))["threat_intel"]
    assert [a["confidence"] for a in t["exploit_attempts"]] == ["low", "high"] and "not in the inventory" in t["exploit_attempts"][0]["confidence_reason"]
    assert [c["confidence"] for c in t["correlations"]] == ["high", "medium"]
    assert [i["confidence"] for i in t["ioc_hits"]] == ["medium", "low"]


def test_uncertain_items_unpinned_packages_and_inferred_success_are_listed():
    events = [{"id": "LOG-9", "type": "Web attack: SQL injection", "possibly_successful": True}]
    inv = [{"name": "requests", "version": "unpinned", "ecosystem": "pypi"}, {"name": "flask", "version": "1.0", "ecosystem": "pypi"}]
    patch = review(state_with(cves=[{"ecosystem": "os"}], events=events, inventory=inv))
    kinds = {u["kind"] for u in patch["critic"]["uncertain"]}
    assert {"CVE match", "Log detection"} <= kinds
    assert patch["critic"]["unverified"] == [{"component": "requests", "ecosystem": "pypi", "reason": "version is not pinned, so known CVEs could not be matched"}]
    assert patch["critic"]["counts"] == {"high": 0, "medium": 1, "low": 0}


def test_caveats_are_added_to_the_models_text_so_it_cannot_sound_surer_than_the_data():
    patch = review(state_with(cves=[{"ecosystem": "os"}], inventory=[{"name": "x", "version": "unpinned", "ecosystem": "pypi"}]))
    for text in (patch["threat_intel"]["summary"], patch["incident_plan"]["briefing"]):
        assert text.startswith(("Model summary.", "Model briefing.")) and "Review notes (uncertainty critic)" in text
        assert "1 of 1 CVE matches are lower confidence" in text and "1 dependencies are not pinned" in text
    no_briefing = review({**state_with(cves=[{"ecosystem": "os"}]), "incident_plan": {"incidents": [], "steps": []}})
    assert no_briefing["incident_plan"]["briefing"].startswith("**Review notes")


def test_the_review_changes_nothing_in_place_and_is_repeatable():
    state = state_with(cves=[{"ecosystem": "os"}])
    before = copy.deepcopy(state)
    first = review(state)
    assert state == before and "confidence" not in state["threat_intel"]["cve_matches"][0]
    again = review(state)
    assert {k: v for k, v in first.items() if k != "trace"} == {k: v for k, v in again.items() if k != "trace"}


def test_the_critic_on_the_real_demo_marks_the_uncertain_things():
    from secagents.graph import run

    state = dict(run())
    assert state["critic"]["reviewed"] > 0 and state["critic"]["counts"]["high"] > 0
    attempt = next(a for a in state["threat_intel"]["exploit_attempts"] if a["cve"] == "CVE-2021-44228")
    assert attempt["confidence"] == "low"  # Log4Shell was attempted, but log4j is not installed
    assert all("confidence" in m for m in state["threat_intel"]["cve_matches"])
    assert "Review notes" in state["threat_intel"]["summary"]


# ---- correlation ----------------------------------------------------------------------------------------------------------------
def test_correlation_joins_events_weaknesses_cves_techniques_and_control_gaps():
    state = {
        "log_findings": {"events": [
            {"id": "LOG-1", "type": "Successful login after brute force", "severity": "CRITICAL", "src_ip": "203.0.113.45", "mitre": ["T1110.001", "T1078"]},
            {"id": "LOG-2", "type": "Web attack: SQL injection", "severity": "HIGH", "src_ip": "192.0.2.77", "mitre": ["T1190"]},
            {"id": "LOG-3", "type": "Port scan", "severity": "HIGH", "src_ip": "198.51.100.9", "mitre": ["T1046"]}]},
        "vuln_findings": {"findings": [{"id": "VULN-1"}]},
        "threat_intel": {"correlations": [{"event": "LOG-2", "finding": "VULN-1", "verdict": "Likely exploited: attack returned HTTP 2xx"}],
                         "exploit_attempts": [{"event": "LOG-2", "cve": "CVE-2021-44228", "component": "log4j-core", "component_present": True}],
                         "cve_matches": [{"cve": "CVE-2021-44228", "package": "log4j-core"}], "ioc_hits": [{"indicator": "203.0.113.45"}]},
        "incident_plan": {"incidents": [{"id": "INC-01", "title": "Host compromise", "priority": "P1", "events": ["LOG-1", "LOG-2"]}]},
        "compliance": {"gaps": [{"check": "account_lockout", "title": "Account lockout"}, {"check": "secure_code", "title": "Secure code"}, {"check": "patch_sla", "title": "Patch SLA"}]},
    }
    out = correlate(state)
    [link] = out["links"]
    assert link["events"] == ["LOG-1", "LOG-2"] and link["attackers"] == ["192.0.2.77", "203.0.113.45"]
    assert link["weaknesses"] == ["VULN-1"] and link["confirmed_exploited"] == ["VULN-1"] and link["cves"] == ["CVE-2021-44228"]
    assert link["techniques"] == ["T1078", "T1110.001", "T1190"] and link["known_bad"] == ["203.0.113.45"]
    assert set(link["control_gaps"]) == {"Account lockout", "Secure code"}  # Patch SLA is not linked: no Log4Shell event in this incident
    assert out["unlinked_events"] == [{"id": "LOG-3", "type": "Port scan", "severity": "HIGH"}]
    assert "INC-01: 2 log events" in out["summary"] and "not part of any incident" in out["summary"]


def test_correlation_with_nothing_to_correlate_is_empty_not_an_error():
    assert correlate({})["summary"] == "No incidents to correlate." and correlate({})["links"] == []


# ---- the human gate ---------------------------------------------------------------------------------------------------------------
PLAN = {"steps": [
    {"step": 1, "incident": "INC-01", "phase": "Detection & Analysis", "action": "Confirm scope", "command": "last -F", "priority": "P1"},
    {"step": 2, "incident": "INC-01", "phase": "Containment", "action": "Block attacker", "command": "sudo ufw insert 1 deny from 203.0.113.45", "priority": "P1"},
    {"step": 3, "incident": "INC-01", "phase": "Eradication", "action": "Remove backdoor", "command": "sudo userdel -r toor", "priority": "P1"},
    {"step": 4, "incident": "INC-01", "phase": "Containment", "action": "Notify the on-call lead", "command": "", "priority": "P2"},
    {"step": 5, "incident": "ALL", "phase": "Post-Incident", "action": "Post-mortem", "command": "", "priority": "P3"}]}


def test_only_containment_and_eradication_steps_need_approval_and_nothing_is_executed():
    gate = approvals.propose(PLAN)
    assert [a["id"] for a in gate["actions"]] == ["S2", "S3", "S4"] and gate["executes"] is False
    assert gate["status"] == "pending_approval" and all(a["status"] == "pending" and a["approved_command"] == "" for a in gate["actions"])
    assert approvals.propose({"steps": [PLAN["steps"][0]]})["status"] == "none"


def test_approving_editing_and_rejecting_are_recorded_with_who_and_when():
    gate = approvals.propose(PLAN)
    out = approvals.decide(gate, {"S2": {"decision": "approve", "command": "sudo ufw insert 1 deny from 203.0.113.45 comment 'INC-01'"},
                                  "S3": {"decision": "reject"}}, "alice@example.com", now="2026-10-23T10:00:00Z")
    by = {a["id"]: a for a in out["actions"]}
    assert by["S2"]["status"] == "approved" and by["S2"]["approved_command"].endswith("comment 'INC-01'")
    assert (by["S2"]["decided_by"], by["S2"]["decided_at"]) == ("alice@example.com", "2026-10-23T10:00:00Z")
    assert by["S3"]["status"] == "rejected" and by["S3"]["approved_command"] == "" and by["S4"]["status"] == "pending"
    assert out["status"] == "in_review" and approvals.counts(out) == {"pending": 1, "approved": 1, "rejected": 1}
    first = out["audit"][0]
    assert (first["by"], first["action"], first["decision"], first["from"]) == ("alice@example.com", "S2", "approved", "pending")
    assert first["command_before"] == "sudo ufw insert 1 deny from 203.0.113.45" and "comment" in first["command_after"]
    assert gate["actions"][0]["status"] == "pending" and gate["audit"] == []  # the input is not modified


def test_unchanged_decisions_are_not_audited_again_and_a_change_of_mind_is():
    gate = approvals.propose(PLAN)
    one = approvals.decide(gate, {"S2": {"decision": "approve"}}, "a", now="t1")
    same = approvals.decide(one, {"S2": {"decision": "approve", "command": PLAN["steps"][1]["command"]}}, "b", now="t2")
    assert len(same["audit"]) == 1 and same["actions"][0]["decided_by"] == "a"
    changed = approvals.decide(same, {"S2": {"decision": "reject"}}, "b", now="t3")
    assert [e["decision"] for e in changed["audit"]] == ["approved", "rejected"] and changed["actions"][0]["approved_command"] == ""
    back = approvals.decide(changed, {"S2": {"decision": "pending"}}, "b", now="t4")
    assert back["actions"][0]["status"] == "pending" and back["actions"][0]["decided_by"] == ""


def test_overall_status_follows_the_decisions():
    gate = approvals.propose(PLAN)
    everything = {a["id"]: {"decision": "approve"} for a in gate["actions"]}
    assert approvals.decide(gate, everything, "a")["status"] == "approved"
    assert approvals.decide(gate, {k: {"decision": "reject"} for k in everything}, "a")["status"] == "rejected"
    assert approvals.decide(gate, {**everything, "S3": {"decision": "reject"}}, "a")["status"] == "decided"


def test_bad_decisions_are_refused_without_changing_anything():
    gate = approvals.propose(PLAN)
    with pytest.raises(ValueError, match="unknown action"):
        approvals.decide(gate, {"S99": {"decision": "approve"}}, "a")
    with pytest.raises(ValueError, match="must be one of"):
        approvals.decide(gate, {"S2": {"decision": "yes please"}}, "a")
    for bad in ("ls\nrm -rf /", "ls\rid", "ls\x00"):
        with pytest.raises(ValueError, match="single line"):
            approvals.decide(gate, {"S2": {"decision": "approve", "command": bad}}, "a")  # hidden second commands cannot be smuggled into the runbook
    assert gate["status"] == "pending_approval" and gate["audit"] == []


def test_the_runbook_contains_only_approved_commands_and_says_it_was_not_run():
    gate = approvals.decide(approvals.propose(PLAN), {"S2": {"decision": "approve", "command": "sudo ufw deny from 203.0.113.45"},
                                                      "S3": {"decision": "reject"}, "S4": {"decision": "approve"}}, "alice@example.com", now="2026-10-23T10:00:00Z")
    script = approvals.runbook(gate, "INC-01 containment", who="alice@example.com")
    assert script.startswith("#!/usr/bin/env bash") and "did NOT run any of this" in script
    assert "sudo ufw deny from 203.0.113.45" in script and "userdel" not in script  # the rejected step is not in it
    assert "[S2] INC-01 · Containment · P1: Block attacker" in script and "approved by alice@example.com at 2026-10-23T10:00:00Z" in script
    assert "(no command: carry out this step manually)" in script  # approved step S4 has no command
    assert "No actions have been approved yet." in approvals.runbook(approvals.propose(PLAN))


def test_the_real_demo_proposes_actions_and_executes_nothing():
    from secagents.graph import run

    gate = dict(run())["containment_gate"]
    assert gate["status"] == "pending_approval" and gate["executes"] is False and gate["actions"]
    assert {a["phase"] for a in gate["actions"]} <= {"Containment", "Eradication"} and all(a["status"] == "pending" for a in gate["actions"])


# ---- reports show the new steps ----------------------------------------------------------------------------------------------
@pytest.fixture(scope="module")
def demo():
    from secagents.graph import run

    return dict(run())


def test_the_markdown_report_has_confidence_correlation_gate_and_review(demo):
    md = demo["report_md"]
    assert "**Confidence:** " in md and "see section 6" in md
    assert "| Confidence | Source | CISA KEV |" in md  # the CVE table says how sure each match is and where it came from
    assert "### How the findings connect (supervisor correlation)" in md and "### Containment approval (human gate)" in md
    assert "awaiting a person's approval" in md and "executes none of them" in md
    assert "## 6. Uncertainty review: what to treat with caution" in md and "Review notes (uncertainty critic)" in md


def test_excel_word_and_pdf_carry_the_same_additions(demo):
    import io

    from docx import Document
    from openpyxl import load_workbook

    from secagents import exporters

    wb = load_workbook(io.BytesIO(exporters.to_xlsx(demo, {})))
    assert {"Correlation", "Containment", "Review", "Unverified", "Approval audit"} <= set(wb.sheetnames)
    headers = [c.value for c in wb["CVEs"][1]]
    assert headers[-3:] == ["Confidence", "Data source", "CISA KEV"] and headers[4] == "CVSS"  # new columns at the end: existing ones did not move
    assert wb["Containment"].max_row - 1 == len(demo["containment_gate"]["actions"])
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=1)}
    assert "Findings reviewed by the critic (high / medium / low confidence)" in summary
    assert summary["Containment actions (this tool executes none of them)"].startswith(f"{len(demo['containment_gate']['actions'])} awaiting approval")
    doc = Document(io.BytesIO(exporters.to_docx(demo, {})))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "6. Uncertainty review: what to treat with caution" in heads and any("human approval gate" in h for h in heads)
    assert exporters.to_pdf(demo, {})[:5] == b"%PDF-"


def test_decisions_are_stored_in_the_saved_scan_and_survive_a_reload(demo, tmp_path, monkeypatch):
    import copy
    from dataclasses import replace

    from secagents import history
    from secagents.config import settings

    monkeypatch.setattr(history, "settings", replace(settings, history_dir=tmp_path))
    monkeypatch.setattr(history, "_remember_incidents", lambda *a: None)
    sid = history.save(copy.deepcopy(demo), "demo", "test")
    gate = approvals.decide(demo["containment_gate"], {demo["containment_gate"]["actions"][0]["id"]: {"decision": "approve"}}, "alice", now="2026-10-23T10:00:00Z")
    history.set_gate(sid, gate)
    stored = history.load(sid)["result"]["containment_gate"]
    assert stored["actions"][0]["status"] == "approved" and stored["audit"][0]["by"] == "alice" and stored["status"] == "in_review"
